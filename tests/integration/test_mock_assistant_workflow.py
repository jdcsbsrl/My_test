"""Mock-only acceptance of real generation policy, audit, registry and reports."""

from copy import deepcopy

import pytest
from openpyxl import load_workbook

from fixtures.mock_order_scenario import MockKnowledgeAPI, REQUIREMENT_ID, RULES, candidate, query_response
from modules.auto_test.core.regression_checks import execute_sales_queries
from modules.trae_test.orchestrator.audit_gateway import AuditGateway
from modules.trae_test.utils.case_generation_service import CaseGenerationService
from modules.trae_test.utils.case_registry import CaseRegistry
from modules.trae_test.utils.excel_generator import ExcelGenerator
from modules.trae_test.utils.template_builder import ALL_FIELDS
from tools.report_generator import TestReportGenerator as Report


@pytest.fixture
def service(monkeypatch):
    from modules.trae_test.utils import dir_validator

    monkeypatch.setattr(dir_validator, "_load_module_hierarchy", lambda: {"销售": {"订单处理": ["销售订单"]}})
    return CaseGenerationService(retriever=MockKnowledgeAPI(), gateway=AuditGateway())


def test_generate_export_and_reuse_without_fixed_field_confirmation(service, tmp_path):
    registry = CaseRegistry(tmp_path / "mock_cases.sqlite3")
    cases = []
    for kind in RULES:
        result = service.generate(RULES[kind], REQUIREMENT_ID, candidate=candidate(kind))
        assert result["status"] == "ready", result
        assert result["questions"] == []
        case = result["cases"][0]
        assert case["用例状态"] == "正常"
        assert case["预期结果"] == candidate(kind)["预期结果"]
        assert registry.save(case, REQUIREMENT_ID) == registry.save(case, REQUIREMENT_ID)
        cases.append(case)
    assert len(registry.list_cases()) == 7
    path = ExcelGenerator.generate_excel(
        cases, requirement_name="Mock订单查询", output_path=str(tmp_path / "mock_cases.xlsx")
    )
    workbook = load_workbook(path)
    try:
        sheet = workbook.active
        assert [cell.value for cell in sheet[1]] == ALL_FIELDS
        assert sheet.max_row == 8 and sheet.max_column == 15
    finally:
        workbook.close()


@pytest.mark.parametrize("fault", ["empty", "business_error", "ignored_filters"])
def test_fault_injection_prevents_false_success(fault):
    class Facade:
        def query_orders(self, page_num=1, page_size=10, **filters):
            return query_response(
                {"pageNum": page_num, "pageSize": page_size, **filters},
                empty=fault == "empty",
                business_code=500 if fault == "business_error" else 200,
                ignore_filters=fault == "ignored_filters",
            )

    report = Report()
    execute_sales_queries(Facade(), report)
    assert report.exit_code == (2 if fault == "empty" else 1)


def test_missing_steps_are_generation_defect_not_user_confirmation(service):
    draft = candidate("exact")
    draft["用例步骤"] = ""
    result = service.generate(RULES["exact"], REQUIREMENT_ID, candidate=draft)
    assert result["status"] == "generation_failed"
    assert result["questions"] == []


def test_semantic_audit_boundary_is_explicit(service):
    # Reviewed finite source wording must reject the previously accepted contradiction.
    draft = deepcopy(candidate("exact"))
    draft["预期结果"] = "1. HTTP状态和业务状态码均为200\n2. 精确查询返回所有24条订单记录"
    result = service.generate(RULES["exact"], REQUIREMENT_ID, candidate=draft)
    assert result["status"] == "generation_failed"
    response = query_response({"orderNo": "MOCK-ORDER-001", "pageSize": 10})
    assert len(response.json()["data"]["tableDataInfo"]["rows"]) == 1
    assert any(issue["code"] == "TC_RULE_CONTRACT" for issue in result["audit"]["errors"])
